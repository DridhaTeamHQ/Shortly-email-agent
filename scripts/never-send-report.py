"""Excel of addresses added to the never-send list, with the status of each.

    python -I scripts/never-send-report.py                 # today (IST)
    python -I scripts/never-send-report.py --date 2026-10-13
    python -I scripts/never-send-report.py --from 2026-10-07 --to 2026-10-13
    python -I scripts/never-send-report.py --all

Writes ../reports/never-send-<period>.xlsx. Reads SUPABASE_URL and
SUPABASE_SERVICE_ROLE_KEY from .env (never printed).
"""

import argparse
import json
import re
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
IST = timezone(timedelta(hours=5, minutes=30))
STATUS = {  # keep in sync with supabase/functions/_shared/never-send-status.ts
    "precheck_domain": "precheck-domain (not an allowed domain)",
    "precheck_validity": "precheck-validity",
    "ses_medium_confidence": "medium-overall-confidence (SES)",
    "ses_low_confidence": "low-overall-confidence (SES)",
    "ses_auto_validation": "ses-auto-validation (blocked before sending)",
    "hard_bounce": "hard-bounce",
    "repeated_soft_bounce": "repeated-soft-bounce",
    "complaint": "spam-complaint",
    "unsubscribed": "unsubscribed",
    "manual": "manual",
}
WHY = {
    "precheck-domain (not an allowed domain)": "Not an @gmail.com address. Only @gmail.com is accepted.",
    "precheck-validity": "Failed our pre-check (shared mailbox, impossible Gmail name, random-looking name or misspelled domain).",
    "medium-overall-confidence (SES)": "SES rated the address MEDIUM - not confident the mailbox exists. Only HIGH is mailed.",
    "low-overall-confidence (SES)": "SES rated the address LOW - very likely does not exist.",
    "ses-auto-validation (blocked before sending)": "SES Auto Validation blocked it before sending (below HIGH).",
    "hard-bounce": "The mailbox rejected our email - the account does not exist or is closed.",
    "repeated-soft-bounce": "Bounced 5+ times in 14 days (e.g. mailbox full) with nothing delivered.",
    "spam-complaint": "The person marked our email as spam.",
    "unsubscribed": "The person unsubscribed.",
    "manual": "Added by the team.",
}


def status_of(row):
    # Entries written before the 2026-10-13 status update carry the status in
    # their detail text; show them with the status they will have.
    if row["reason"] == "manual" and (row.get("detail") or "").startswith("precheck-domain"):
        return STATUS["precheck_domain"]
    if row["reason"] == "hard_bounce" and row.get("detail") == "EmailValidationSuppressed":
        return STATUS["ses_auto_validation"]
    return STATUS.get(row["reason"], row["reason"])


def load_env():
    env = {}
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\s*([A-Z0-9_]+)\s*=\s*(.*?)\s*$", line)
        if m:
            env[m.group(1)] = m.group(2).strip('"').strip("'")
    return env


def ist_start(day):
    return datetime.fromisoformat(f"{day}T00:00:00+05:30").astimezone(timezone.utc)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--from", dest="start")
    ap.add_argument("--to", dest="end")
    ap.add_argument("--all", action="store_true")
    a = ap.parse_args()
    today = datetime.now(IST).date().isoformat()
    start, end = (None, None) if a.all else (a.date or a.start or today, a.date or a.end or a.start or today)
    label = "all" if a.all else (start if start == end else f"{start}_to_{end}")

    env = load_env()
    base, key = env["SUPABASE_URL"].rstrip("/") + "/rest/v1", env["SUPABASE_SERVICE_ROLE_KEY"]
    filters = ""
    if start:
        filters += "&created_at=gte." + quote(ist_start(start).isoformat())
        filters += "&created_at=lt." + quote((ist_start(end) + timedelta(days=1)).isoformat())
    rows, offset = [], 0
    while True:
        req = Request(f"{base}/email_suppressions?select=email,reason,detail,source,created_at&order=created_at.desc{filters}",
                      headers={"apikey": key, "Authorization": f"Bearer {key}", "Range": f"{offset}-{offset + 999}"})
        page = json.loads(urlopen(req, timeout=60).read())
        rows += page
        if len(page) < 1000:
            break
        offset += 1000

    wb = Workbook()
    fill, font = PatternFill("solid", fgColor="3979FF"), Font(bold=True, color="FFFFFF")
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Never-send list", "all time" if a.all else (start if start == end else f"{start} to {end}") + " (IST)"])
    ws.append(["Generated (IST)", datetime.now(IST).strftime("%Y-%m-%d %H:%M")])
    ws.append(["Total addresses", len(rows)])
    ws.append(["", ""])
    for status, n in Counter(status_of(r) for r in rows).most_common():
        ws.append([status, n])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 46, 30
    ws["A1"].font = Font(bold=True, size=13)

    ls = wb.create_sheet("Never-send")
    headers = ["Email", "Status", "Why it is on the never-send list", "Detail", "Added (IST)", "Source"]
    ls.append(headers)
    for r in rows:
        added = datetime.fromisoformat(r["created_at"].replace("Z", "+00:00")).astimezone(IST).strftime("%Y-%m-%d %H:%M")
        status = status_of(r)
        ls.append([r["email"], status, WHY.get(status, ""), r.get("detail") or "", added, r.get("source") or ""])
    for c, w in enumerate([36, 40, 70, 50, 17, 30], start=1):
        cell = ls.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment = fill, font, Alignment(vertical="center")
        ls.column_dimensions[get_column_letter(c)].width = w
    ls.freeze_panes = "A2"
    ls.auto_filter.ref = ls.dimensions

    out = ROOT.parent / "reports" / f"never-send-{label}.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(json.dumps({"rows": len(rows), "by_status": Counter(status_of(r) for r in rows), "file": str(out)}, indent=1))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error)}), file=sys.stderr)
        sys.exit(1)
